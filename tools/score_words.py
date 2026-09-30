"""Assign reproducible 0-100 lexical difficulty scores to active Italian words."""

from __future__ import annotations

import hashlib
import io
import os
import sys
import urllib.request
from datetime import datetime, timezone

import numpy as np
import openpyxl
import psycopg
from wordfreq import zipf_frequency


DATA_URL = os.getenv("ITAOA_URL", "https://osf.io/download/3nvh6/")
DATA_SHA256 = "026bcdba1c9544167829f19d7b0cd53fa21836fc7c24891a92eee75fbaf4aabc"
MODEL_VERSION = "itaoa-wordfreq-ridge-v1"
MIN_AOA = 2.0
MAX_AOA = 18.0


def download_itaoa() -> bytes:
    request = urllib.request.Request(DATA_URL, headers={"User-Agent": "Nodo vocabulary scorer/1.0"})
    with urllib.request.urlopen(request, timeout=60) as response:
        content = response.read()
    if DATA_URL == "https://osf.io/download/3nvh6/":
        digest = hashlib.sha256(content).hexdigest()
        if digest != DATA_SHA256:
            raise RuntimeError(f"ItAoA checksum mismatch: {digest}")
    return content


def load_ratings(content: bytes) -> list[dict[str, object]]:
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.worksheets[1]
    headers = [str(value) for value in next(sheet.iter_rows(values_only=True))]
    rows: list[dict[str, object]] = []
    for values in sheet.iter_rows(values_only=True):
        row = dict(zip(headers, values))
        word = str(row.get("Ita_Word") or "").strip().lower()
        pos = str(row.get("WordClass") or "").strip().lower()
        if word and pos in {"n", "v"} and row.get("M_AoA") is not None:
            rows.append({"word": word, "pos": pos, "aoa": float(row["M_AoA"])})
    if len(rows) < 1_000:
        raise RuntimeError(f"ItAoA parsing returned only {len(rows)} noun/verb ratings")
    return rows


def approximate_syllables(word: str) -> int:
    vowels = "aeiouàèéìòóù"
    groups = 0
    previous_vowel = False
    for character in word:
        current_vowel = character in vowels
        if current_vowel and not previous_vowel:
            groups += 1
        previous_vowel = current_vowel
    return max(1, groups)


def features(word: str, pos: str) -> np.ndarray:
    letters = [character for character in word if character.isalpha()]
    vector = np.zeros(135, dtype=float)
    vector[:7] = [
        zipf_frequency(word, "it", wordlist="best", minimum=0.0),
        len(letters),
        approximate_syllables(word),
        sum(character in "aeiouàèéìòóù" for character in letters) / max(1, len(letters)),
        len(set(letters)) / max(1, len(letters)),
        float(pos == "v"),
        float(word.endswith("si")),
    ]
    grams = [word[index:index + size] for size in (2, 3, 4) for index in range(max(0, len(word) - size + 1))]
    for gram in grams:
        bucket = int.from_bytes(hashlib.blake2b(gram.encode("utf-8"), digest_size=2).digest(), "big") % 128
        vector[7 + bucket] += 1.0 / max(1, len(grams))
    return vector


def fit_ridge(x: np.ndarray, y: np.ndarray, penalty: float) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    mean = x.mean(axis=0)
    scale = x.std(axis=0)
    scale[scale < 1e-9] = 1.0
    standardized = (x - mean) / scale
    design = np.column_stack((np.ones(len(x)), standardized))
    regularizer = np.eye(design.shape[1]) * penalty
    regularizer[0, 0] = 0.0
    coefficients = np.linalg.solve(design.T @ design + regularizer, design.T @ y)
    return mean, scale, coefficients


def predict_ridge(x: np.ndarray, model: tuple[np.ndarray, np.ndarray, np.ndarray]) -> np.ndarray:
    mean, scale, coefficients = model
    design = np.column_stack((np.ones(len(x)), (x - mean) / scale))
    return design @ coefficients


def cross_validated_model(x: np.ndarray, y: np.ndarray) -> tuple[tuple[np.ndarray, np.ndarray, np.ndarray], float, float]:
    random = np.random.default_rng(20260929)
    indices = random.permutation(len(x))
    folds = np.array_split(indices, 5)
    best_penalty = 0.0
    best_mae = float("inf")
    for penalty in (0.1, 1.0, 10.0, 100.0):
        errors = []
        for validation in folds:
            training = np.setdiff1d(indices, validation, assume_unique=True)
            prediction = predict_ridge(x[validation], fit_ridge(x[training], y[training], penalty))
            errors.extend(np.abs(prediction - y[validation]))
        mae = float(np.mean(errors))
        if mae < best_mae:
            best_penalty, best_mae = penalty, mae
    return fit_ridge(x, y, best_penalty), best_mae, best_penalty


def score_from_aoa(aoa: float) -> float:
    return round(float(np.clip((aoa - MIN_AOA) / (MAX_AOA - MIN_AOA) * 100, 0, 100)), 2)


def main() -> None:
    database_url = os.getenv("DATABASE_URL")
    if not database_url:
        raise RuntimeError("DATABASE_URL is required")

    print("Downloading and validating ItAoA...", flush=True)
    ratings = load_ratings(download_itaoa())
    print(f"Loaded {len(ratings)} noun/verb ratings; training model...", flush=True)
    x = np.vstack([features(str(row["word"]), str(row["pos"])) for row in ratings])
    y = np.array([float(row["aoa"]) for row in ratings])
    model, mae, penalty = cross_validated_model(x, y)
    prediction_confidence = round(float(np.clip(1.0 - mae / 8.0, 0.35, 0.80)), 3)
    direct = {(str(row["word"]), str(row["pos"])): float(row["aoa"]) for row in ratings}

    print("Model trained; scoring active database words...", flush=True)
    with psycopg.connect(database_url) as connection:
        with connection.cursor() as cursor:
            cursor.execute("SELECT id,word,part_of_speech FROM words WHERE active ORDER BY id")
            words = cursor.fetchall()
            predicted_features = np.vstack([features(word, pos) for _, word, pos in words])
            predicted_aoa = predict_ridge(predicted_features, model)
            updates = []
            direct_count = 0
            for (word_id, word, pos), estimate, feature_row in zip(words, predicted_aoa, predicted_features):
                measured = direct.get((word, pos))
                if measured is not None:
                    aoa = measured
                    source = "itaoa-human"
                    confidence = 0.95
                    direct_count += 1
                else:
                    aoa = round(float(np.clip(estimate, MIN_AOA, MAX_AOA)), 2)
                    source = "itaoa-ridge"
                    confidence = prediction_confidence
                updates.append((
                    score_from_aoa(aoa), source, confidence, aoa,
                    round(float(feature_row[0]), 2), MODEL_VERSION,
                    datetime.now(timezone.utc), word_id,
                ))
            cursor.executemany(
                """
                UPDATE words SET
                    difficulty_score=%s,difficulty_source=%s,difficulty_confidence=%s,
                    age_of_acquisition=%s,frequency_zipf=%s,difficulty_model_version=%s,
                    difficulty_scored_at=%s
                WHERE id=%s
                """,
                updates,
            )
            cursor.execute("SELECT count(*) FROM words WHERE active AND difficulty_score IS NULL")
            missing = cursor.fetchone()[0]
            if missing:
                raise RuntimeError(f"scoring left {missing} active words without a score")
        connection.commit()

    print(f"Scored {len(words)} active words: {direct_count} human-rated, {len(words)-direct_count} predicted")
    print(f"Five-fold prediction MAE: {mae:.2f} acquisition-years; predicted confidence: {prediction_confidence:.3f}")
    print(f"Selected ridge penalty: {penalty:g}")
    print(f"Model version: {MODEL_VERSION}")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"word scoring failed: {error}", file=sys.stderr)
        raise
